"""Read CSV, XLSX, XML, and MetaQuotes HTML statements into a grid or ready rows."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from app.services.imports.headers import detect_header_index

_MAX_ROWS = 20_000
_MAX_COLS = 80
_HTML_MARKERS = ("<html", "<table", "<tr")


class ImportFileError(Exception):
    def __init__(self, message: str):
        self.message = message
        super().__init__(message)


@dataclass
class ParsedFile:
    kind: str
    headers: list[str]
    records: list[tuple[int, dict[str, str]]] = field(default_factory=list)
    prebuilt: list[dict] | None = None


def parse_file(filename: str, payload: bytes) -> ParsedFile:
    if not payload:
        raise ImportFileError("The file is empty.")
    name = filename.lower()
    if name.endswith(".xlsx"):
        return _parse_xlsx(payload)
    if name.endswith(".xml"):
        return _parse_xml(payload)
    if name.endswith(".htm") or name.endswith(".html") or _looks_like_html(payload):
        text = _decode_text(payload)
        if "closed transactions" in text.lower():
            kind = "mt4-html" if "mt4" in text.lower() or "metatrader 4" in text.lower() else "mt5-html"
            rows = _parse_mt_html(text)
            headers = list(rows[0]["raw"].keys()) if rows else []
            return ParsedFile(kind=kind, headers=headers, prebuilt=rows)
        if name.endswith(".htm") or name.endswith(".html"):
            raise ImportFileError("This HTML file is not a MetaTrader closed-transactions statement.")
    if name.endswith(".csv") or _looks_like_text(payload):
        return _parse_csv(payload)
    raise ImportFileError("Upload a CSV, Excel (.xlsx), or XML file.")


def _parse_csv(payload: bytes) -> ParsedFile:
    text = _decode_text(payload)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    grid = _numbered([_trim_row(row) for row in reader])
    return _from_grid(grid, kind="csv")


def _parse_xlsx(payload: bytes) -> ParsedFile:
    if not payload.startswith(b"PK"):
        raise ImportFileError("This Excel file is not a valid workbook.")
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("This Excel file could not be read.") from exc
    try:
        sheet = _pick_sheet(workbook)
        grid = _numbered(
            [_trim_row([_excel_cell(value) for value in row]) for row in sheet.iter_rows(values_only=True)]
        )
    finally:
        workbook.close()
    parsed = _from_grid(grid, kind="xlsx")
    return parsed


def _parse_xml(payload: bytes) -> ParsedFile:
    head = payload[:8000].lower()
    if b"<!doctype" in head or b"<!entity" in head:
        raise ImportFileError("This XML file can't be imported.")
    stripped = payload.lstrip()
    if not stripped.startswith(b"<") and not stripped.startswith(b"\xef\xbb\xbf<"):
        raise ImportFileError("This XML file can't be imported.")
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ImportFileError("This XML file could not be read.") from exc
    spreadsheet = _spreadsheet_grid(root)
    if spreadsheet is not None:
        return _from_grid(spreadsheet, kind="xml")
    records = _record_elements(root)
    if not records:
        raise ImportFileError("No trade rows were found in this XML file.")
    if len(records) > _MAX_ROWS:
        raise ImportFileError("This file has too many rows.")
    headers: list[str] = []
    seen: set[str] = set()
    raw_rows: list[tuple[int, dict[str, str]]] = []
    for index, element in enumerate(records, start=1):
        cells: dict[str, str] = {}
        for key, value in element.attrib.items():
            name = _local(key)
            cells[name] = value.strip()
            if name not in seen:
                headers.append(name)
                seen.add(name)
        for child in list(element):
            name = _local(child.tag)
            if name in cells:
                continue
            cells[name] = (child.text or "").strip()
            if name not in seen:
                headers.append(name)
                seen.add(name)
        if any(cells.values()):
            raw_rows.append((index, cells))
    if not raw_rows:
        raise ImportFileError("No trade rows were found in this XML file.")
    kind = "tradefix-xml" if _local(root.tag) == "tradefix" else "xml"
    return ParsedFile(kind=kind, headers=headers[:_MAX_COLS], records=raw_rows)


def _from_grid(grid: list[tuple[int, list[str]]], *, kind: str) -> ParsedFile:
    usable = [(number, cells) for number, cells in grid if any(cell.strip() for cell in cells)]
    if not usable:
        raise ImportFileError("No trade rows were found in this file.")
    header_at = detect_header_index(usable)
    header_cells = usable[header_at][1]
    headers = _unique_headers(header_cells)
    records: list[tuple[int, dict[str, str]]] = []
    for number, cells in usable[header_at + 1 :]:
        if not any(cell.strip() for cell in cells):
            continue
        raw = {headers[index]: cells[index].strip() if index < len(cells) else "" for index in range(len(headers))}
        if any(raw.values()):
            records.append((number, raw))
        if len(records) > _MAX_ROWS:
            raise ImportFileError("This file has too many rows.")
    if not records:
        raise ImportFileError("No trade rows were found under the header row.")
    detected = kind
    labels = {header.strip().lower() for header in headers}
    if {"symbol", "trade_date", "trade_type", "quantity", "price"}.issubset(labels):
        detected = "zerodha-tradebook-csv"
    return ParsedFile(kind=detected, headers=headers, records=records)


def _parse_mt_html(text: str) -> list[dict]:
    rows: list[dict] = []
    body = text
    marker = re.search(r"closed transactions", text, flags=re.IGNORECASE)
    if marker:
        body = text[marker.start() :]
    for index, match in enumerate(re.findall(r"<tr[^>]*>(.*?)</tr>", body, flags=re.IGNORECASE | re.DOTALL)):
        cells = [
            re.sub(r"<[^>]+>", "", cell).strip()
            for cell in re.findall(r"<td[^>]*>(.*?)</td>", match, flags=re.IGNORECASE | re.DOTALL)
        ]
        if len(cells) < 8 or not cells[0].isdigit():
            continue
        raw = {
            "ticket": cells[0],
            "open_time": cells[1],
            "type": cells[2],
            "size": cells[3],
            "item": cells[4],
            "open_price": cells[5],
            "close_time": cells[8] if len(cells) > 8 else "",
            "close_price": cells[9] if len(cells) > 9 else "",
            "commission": cells[10] if len(cells) > 10 else "",
            "swap": cells[12] if len(cells) > 12 else "",
        }
        errors: list[str] = []
        try:
            opened = _parse_mt_dt(raw["open_time"])
        except ValueError:
            opened = None
            errors.append("open time was not understood")
        side = raw["type"].lower()
        if side not in {"buy", "sell"}:
            continue
        normalized = {
            "symbol": raw["item"],
            "side": side,
            "quantity": raw["size"],
            "price": raw["open_price"],
            "exit_price": raw["close_price"] or None,
            "executed_at": opened.isoformat() if opened else None,
            "closed_at": raw["close_time"] or None,
            "commission": raw["commission"] or None,
            "swap": raw["swap"] or None,
            "external_id": raw["ticket"],
        }
        rows.append(
            {
                "row_number": index + 1,
                "raw": raw,
                "normalized": normalized,
                "errors": errors,
                "status": "attention" if errors else "valid",
            }
        )
    if not rows:
        raise ImportFileError("No closed trades were found in this statement.")
    return rows


def _spreadsheet_grid(root: ET.Element) -> list[tuple[int, list[str]]] | None:
    rows = [element for element in root.iter() if _local(element.tag) == "row"]
    if len(rows) < 2:
        return None
    grid: list[tuple[int, list[str]]] = []
    for index, row in enumerate(rows, start=1):
        cells: list[str] = []
        for cell in list(row):
            if _local(cell.tag) != "cell":
                continue
            data = next((child for child in list(cell) if _local(child.tag) == "data"), None)
            text = (data.text if data is not None else cell.text) or ""
            cells.append(text.strip())
        if cells:
            grid.append((index, cells[:_MAX_COLS]))
    return grid or None


def _record_elements(root: ET.Element) -> list[ET.Element]:
    trades = [element for element in root.iter() if _local(element.tag) == "trade"]
    detailed = [element for element in trades if element.attrib or len(list(element))]
    if detailed:
        return detailed[: _MAX_ROWS + 1]
    counts: dict[str, list[ET.Element]] = {}
    for parent in root.iter():
        children = list(parent)
        if len(children) < 1:
            continue
        names = [_local(child.tag) for child in children]
        if len(set(names)) != 1:
            continue
        if not any(child.attrib or len(list(child)) for child in children):
            continue
        counts.setdefault(names[0], []).extend(children)
    if not counts:
        return []
    return max(counts.values(), key=len)[: _MAX_ROWS + 1]


def _pick_sheet(workbook):
    for name in workbook.sheetnames:
        if name.strip().lower() == "trades":
            return workbook[name]
    for name in workbook.sheetnames:
        if name.strip().lower() != "summary":
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def _unique_headers(cells: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(cells[:_MAX_COLS]):
        label = cell.strip() or f"Column {index + 1}"
        count = seen.get(label, 0) + 1
        seen[label] = count
        headers.append(label if count == 1 else f"{label} {count}")
    while headers and headers[-1].startswith("Column "):
        headers.pop()
    return headers or ["Column 1"]


def _numbered(rows: list[list[str]]) -> list[tuple[int, list[str]]]:
    return [(index, row) for index, row in enumerate(rows, start=1)]


def _trim_row(row: list[object]) -> list[str]:
    cells = ["" if cell is None else str(cell).strip() for cell in row[:_MAX_COLS]]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _excel_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _looks_like_html(payload: bytes) -> bool:
    head = payload[:500].lower()
    return all(marker.encode() in head or marker.encode() in payload[:4000].lower() for marker in _HTML_MARKERS[:1]) and (
        b"<tr" in payload[:8000].lower()
    )


def _looks_like_text(payload: bytes) -> bool:
    if payload.startswith((b"PK", b"MZ", b"\x7fELF")):
        return False
    return b"\x00" not in payload[:4096]


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].strip().lower()


def _parse_mt_dt(value: str) -> datetime:
    text = value.strip()
    for fmt in ("%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    raise ValueError(text)
