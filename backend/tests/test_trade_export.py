from datetime import datetime, timezone
from types import SimpleNamespace
from uuid import uuid4
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from app.schemas.trade_export import TradeExportFilters
from app.services.trade_export import (
    columns_for_export,
    export_filename,
    matches_export_filters,
    render_csv,
    render_xlsx,
    render_xml,
    trade_export_row,
)


def _trade(**overrides):
    base = dict(
        id=uuid4(),
        user_id=uuid4(),
        symbol="AAPL",
        asset_type="stock",
        side="long",
        quantity=10,
        entry_price=100,
        exit_price=110,
        opened_at=datetime(2026, 10, 1, 14, 30, tzinfo=timezone.utc),
        closed_at=datetime(2026, 10, 1, 16, 45, tzinfo=timezone.utc),
        pnl=95.5,
        fees=4.5,
        risk_amount=50,
        setup_tag="Breakout",
        setup_tags=["Breakout"],
        strategy_name="Breakout",
        emotion_tags=["Confident"],
        mood="Calm",
        notes='Said "hold", then\nadded size',
        rules_broken=[],
        screenshot_urls=[],
        status="closed",
        session="NY",
        trade_type="Day",
        option_type=None,
        analysis_timeframe="1H",
        entry_timeframe="5m",
        stop_loss=98,
        profit_target=112,
        rating=4,
        invested_amount=1000,
        entry_condition="Breakout",
        exit_condition="Target",
        commission=1.25,
        swap=0,
        funding=None,
        gross_pnl=100,
        strike_price=None,
        expiry_date=None,
        plan_compliance=None,
        source="manual",
        external_trade_id="BRK-1",
        created_at=datetime(2026, 10, 1, 16, 46, tzinfo=timezone.utc),
        extra={"moods": ["Calm"]},
        account=SimpleNamespace(name="Main"),
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def _filters(**overrides) -> TradeExportFilters:
    return TradeExportFilters(**overrides)


def test_filtered_symbol_search_and_status_match_trade_view_rules():
    aapl = _trade()
    other = _trade(symbol="MSFT", setup_tag="Fade", strategy_name="Fade", setup_tags=["Fade"], notes="other")
    closed_filters = _filters(symbol="aapl", search="breakout", status="closed")
    assert matches_export_filters(aapl, closed_filters, "net")
    assert not matches_export_filters(other, closed_filters, "net")
    assert not matches_export_filters(aapl, _filters(status="open"), "net")
    assert not matches_export_filters(aapl, _filters(search="does-not-match"), "net")


def test_pnl_filter_respects_gross_display_mode():
    loser_net = _trade(pnl=-2, fees=5, invested_amount=100)
    assert matches_export_filters(loser_net, _filters(pnl="loss"), "net")
    assert matches_export_filters(loser_net, _filters(pnl="profit"), "gross")
    assert not matches_export_filters(loser_net, _filters(pnl="loss"), "gross")


def test_csv_quotes_special_characters_and_omits_private_fields():
    row = trade_export_row(_trade(notes='AAPL, "quote" & <tag>\nnext'))
    payload = render_csv([row]).decode("utf-8-sig")
    assert payload.startswith("Trade ID,Account,Symbol")
    assert '"AAPL, ""quote"" & <tag>\nnext"' in payload
    assert "user_id" not in payload.splitlines()[0].lower()
    assert "password" not in payload.splitlines()[0].lower()
    assert "AAPL" in payload
    assert "95.50" in payload


def test_xml_escapes_text_and_uses_trade_fields():
    row = trade_export_row(_trade(notes="Tom & Jerry <plan>"))
    root = ET.fromstring(render_xml([row]))
    notes = root.find("./trades/trade/notes")
    assert notes is not None
    assert notes.text == "Tom & Jerry <plan>"
    assert root.find("./trades/trade/symbol").text == "AAPL"
    assert root.find("./trades/trade/strategy").text == "Breakout"
    assert b"user_id" not in render_xml([row])


def test_xlsx_has_trades_summary_and_formats():
    row = trade_export_row(_trade())
    workbook = load_workbook(io_bytes(render_xlsx([row])))
    assert workbook.sheetnames == ["Trades", "Summary"]
    trades = workbook["Trades"]
    assert trades.freeze_panes == "A2"
    assert trades["C1"].value == "Symbol"
    assert trades["C2"].value == "AAPL"
    assert trades["P2"].value == 95.5
    summary = {workbook["Summary"].cell(i, 1).value: workbook["Summary"].cell(i, 2).value for i in range(2, 12)}
    assert summary["Total Trades"] == 1
    assert summary["Winning Trades"] == 1
    assert summary["Losing Trades"] == 0
    assert summary["Total P&L"] == 95.5
    assert summary["Profit Factor"] in ("", None)


def test_export_routes_require_auth():
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    assert client.post("/api/trades/export/count", json={"scope": "all"}).status_code == 401
    assert client.post("/api/trades/export", json={"format": "csv", "scope": "all"}).status_code == 401


def test_custom_columns_keep_picker_order_and_labels():
    row = trade_export_row(_trade())
    chosen = columns_for_export("custom", ["pnl", "ticker", "notes"])
    payload = render_csv([row], chosen).decode("utf-8-sig")
    assert payload.splitlines()[0] == "Ticker,P&L,Notes"
    assert "AAPL" in payload
    assert "95.50" in payload
    assert columns_for_export("filtered", ["ticker"])[0][0] == "Trade ID"


def test_filename_uses_scope_and_date():
    assert export_filename("filtered", "csv", datetime(2026, 10, 8).date()) == "tradefix_trades_filtered_2026-10-08.csv"
    assert export_filename("all", "xlsx", datetime(2026, 10, 8).date()) == "tradefix_trades_2026-10-08.xlsx"
    assert export_filename("all", "xml", datetime(2026, 10, 8).date()) == "tradefix_trades_2026-10-08.xml"
    assert export_filename("custom", "csv", datetime(2026, 10, 8).date()) == "tradefix_trades_custom_2026-10-08.csv"


def io_bytes(payload: bytes):
    import io

    return io.BytesIO(payload)
